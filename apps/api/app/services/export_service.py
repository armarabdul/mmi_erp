import io
import csv
from typing import List, Dict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
import docx

class ExportService:
    @staticmethod
    def export_excel(title: str, columns: List[str], headers: Dict[str, str], data: List[Dict[str, Any]]) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = title[:30]

        # Brand header
        ws.merge_cells("A1:G1")
        title_cell = ws["A1"]
        title_cell.value = f"MMI AI Analytics — {title}"
        title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 32

        # Column headers
        header_labels = [headers.get(c, c) for c in columns]
        ws.append([]) # row 2
        ws.append(header_labels) # row 3
        
        header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for col_idx in range(1, len(columns) + 1):
            cell = ws.cell(row=3, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[3].height = 24

        # Thin border
        thin = Side(border_style="thin", color="E2E8F0")
        border = Border(top=thin, left=thin, right=thin, bottom=thin)

        # Rows
        row_idx = 4
        for row in data:
            row_vals = [row.get(c, "") for c in columns]
            ws.append(row_vals)
            current_row = ws[row_idx]
            for c_idx, cell in enumerate(current_row):
                cell.border = border
                val = cell.value
                if isinstance(val, (int, float)):
                    cell.alignment = Alignment(horizontal="right")
                    if isinstance(val, float):
                        cell.number_format = "#,##0.00"
                else:
                    cell.alignment = Alignment(horizontal="left")
            row_idx += 1

        # Adjust column widths
        for col_idx in range(1, len(columns) + 1):
            col_key = columns[col_idx - 1]
            header_text = str(headers.get(col_key, col_key))
            max_len = max(len(header_text), 10)
            for r in data:
                val_str = str(r.get(col_key, ""))
                if len(val_str) > max_len:
                    max_len = min(len(val_str), 40)
            col_letter = ws.cell(row=3, column=col_idx).column_letter
            ws.column_dimensions[col_letter].width = max_len + 4

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @staticmethod
    def export_pdf(title: str, columns: List[str], headers: Dict[str, str], data: List[Dict[str, Any]]) -> bytes:
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=landscape(letter),
            rightMargin=20,
            leftMargin=20,
            topMargin=20,
            bottomMargin=20,
        )
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            name="TitleStyle",
            parent=styles["Heading1"],
            fontSize=16,
            textColor=colors.HexColor("#1E3A8A"),
            spaceAfter=12,
        )
        elements.append(Paragraph(f"<b>MMI AI Analytics</b> — {title}", title_style))
        elements.append(Paragraph(f"Generated from Demo ERP Analytics Database | Status: Verified", styles["Normal"]))
        elements.append(Spacer(1, 14))

        header_labels = [headers.get(c, c) for c in columns]
        table_rows = [header_labels]
        
        # Limit PDF rows to 100 for clean pagination
        for r in data[:100]:
            table_rows.append([str(r.get(c, "")) for c in columns])

        col_count = len(columns)
        available_width = 750
        col_width = available_width / col_count

        t = Table(table_rows, colWidths=[col_width] * col_count, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
            ('TOPPADDING', (0, 0), (-1, 0), 6),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ]))
        elements.append(t)
        doc.build(elements)
        return output.getvalue()

    @staticmethod
    def export_word(title: str, columns: List[str], headers: Dict[str, str], data: List[Dict[str, Any]]) -> bytes:
        doc = docx.Document()
        doc.add_heading(f"MMI AI Analytics — {title}", level=1)
        doc.add_paragraph("Enterprise Analytical Report generated from MMI ERP Database.")

        table = doc.add_table(rows=1, cols=len(columns))
        table.style = "Light Shading Accent 1"
        hdr_cells = table.rows[0].cells
        for idx, col in enumerate(columns):
            hdr_cells[idx].text = headers.get(col, col)

        for row in data[:200]:
            row_cells = table.add_row().cells
            for idx, col in enumerate(columns):
                row_cells[idx].text = str(row.get(col, ""))

        output = io.BytesIO()
        doc.save(output)
        return output.getvalue()

    @staticmethod
    def export_csv(columns: List[str], headers: Dict[str, str], data: List[Dict[str, Any]]) -> bytes:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([headers.get(c, c) for c in columns])
        for row in data:
            writer.writerow([row.get(c, "") for c in columns])
        return output.getvalue().encode("utf-8-sig")

export_service = ExportService()
